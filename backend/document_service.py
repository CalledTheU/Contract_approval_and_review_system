# Author: WangLei
# Email: WangLei1578@outlook.com
# Date: 2026-09-28

import os
import re
import subprocess
from pathlib import Path

import pymupdf
from docx import Document


def extract(path, suffix):
    if suffix == ".docx":
        doc = Document(path)
        chunks = [paragraph.text for paragraph in doc.paragraphs if paragraph.text.strip()]
        chunks.extend(" | ".join(cell.text for cell in row.cells) for table in doc.tables for row in table.rows)
        return "\n".join(chunks), 1
    if suffix == ".pdf":
        doc = pymupdf.open(path)
        pages = [page.get_text() for page in doc]
        if not any(text.strip() for text in pages):
            exe = os.getenv("TESSERACT_CMD", "tesseract")
            try:
                pages = [subprocess.run([exe, "stdin", "stdout", "-l", os.getenv("OCR_LANG", "chi_sim+eng")], input=page.get_pixmap(dpi=220).tobytes("png"), capture_output=True, timeout=90, check=True).stdout.decode("utf-8", errors="replace") for page in doc]
            except (OSError, subprocess.SubprocessError) as exc:
                raise ValueError(f"扫描版PDF的OCR失败：{exc}; 请安装Tesseract及中文语言包后重试") from exc
            if not any(text.strip() for text in pages):
                raise ValueError("扫描版PDF未识别到正文，请检查扫描质量")
        return "\f".join(pages), len(pages)
    if suffix in {".png", ".jpg", ".jpeg", ".tif", ".tiff"}:
        exe = os.getenv("TESSERACT_CMD", "tesseract")
        try:
            result = subprocess.run([exe, str(path), "stdout", "-l", os.getenv("OCR_LANG", "chi_sim+eng")], capture_output=True, text=True, timeout=90, check=True)
            return result.stdout, 1
        except (OSError, subprocess.SubprocessError) as exc:
            raise ValueError(f"OCR失败：{exc}; 请安装Tesseract及中文语言包后重试") from exc
    raise ValueError("仅支持 DOCX、PDF 和常见图片格式")


def metadata(text):
    amount = re.search(r"(?:人民币|RMB|¥|￥)?\s*([\d,]+(?:\.\d{1,2})?)\s*(万元|万|元)", text)
    parties = re.findall(r"(?:甲方|乙方|买方|卖方|采购方|供应商)\s*[:：]?\s*([^\s，,。；;]{2,40})", text)
    number = re.search(r"(?:合同编号|编号)\s*[:：]?\s*([^\s，,。；;]{2,40})", text)
    period = re.search(r"(?:履行期限|服务期限|合同期限)\s*[:：]?\s*([^\n。；;]{2,60})", text)
    return {
        "amount": amount.group(0).strip() if amount else "未识别",
        "parties": list(dict.fromkeys(parties)),
        "contract_number": number.group(1) if number else "未识别",
        "period": period.group(1).strip() if period else "未识别",
    }
