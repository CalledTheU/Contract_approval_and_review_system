# Author: WangLei
# Email: WangLei1578@outlook.com
# Date: 2026-09-28

import io
import os
import unittest
from unittest.mock import patch

import pymupdf
from docx import Document
from fastapi.testclient import TestClient

import backend.storage as storage
from app import app


class ContractFlowTest(unittest.TestCase):
    def setUp(self):
        self.llm_patch = patch.dict(os.environ, {"LLM_REVIEW_ENABLED": "0"})
        self.llm_patch.start()
        self.uploads = []
        self.db_patch = patch.object(storage, "DB", storage.DATA / "test-contracts.db")
        self.db_patch.start()
        storage.init_db()
        self.client = TestClient(app)
        login = self.client.post("/api/auth/login", json={"username": "legal", "password": "Demo123!"})
        self.assertEqual(login.status_code, 200, login.text)
        self.client.headers.update({"Authorization": f"Bearer {login.json()['token']}"})

    def tearDown(self):
        self.client.close()
        self.db_patch.stop()
        self.llm_patch.stop()
        for path in self.uploads:
            path.unlink(missing_ok=True)
        for suffix in ("", "-wal", "-shm"):
            path = storage.DATA / f"test-contracts.db{suffix}"
            path.unlink(missing_ok=True)

    def test_upload_review_comment_writeback_and_report(self):
        doc = Document()
        doc.add_paragraph("合同编号：TEST-01")
        doc.add_paragraph("乙方完成软件交付后，甲方应支付全部合同价款。")
        doc.add_paragraph("本项目开发成果及相关知识产权均归乙方所有。")
        body = io.BytesIO()
        doc.save(body)
        uploaded = self.client.post("/api/contracts/upload", files={"file": ("risk.docx", body.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")})
        self.assertEqual(uploaded.status_code, 200, uploaded.text)
        result = uploaded.json()
        task_id = result["task"]["id"]
        self.uploads.append(storage.UPLOADS / f"{task_id}.docx")
        self.assertEqual(result["task"]["status"], "completed")
        self.assertTrue(any(risk["level"] == "HIGH" for risk in result["risks"]))
        self.assertTrue(all(risk["original_text"] in result["task"]["text"] for risk in result["risks"]))

        risk = result["risks"][0]
        updated = self.client.put(f"/api/risks/{risk['id']}", json={"accepted": True, "suggested_text": "人工确认的修改条款"})
        self.assertEqual(updated.json()["suggested_text"], "人工确认的修改条款")
        self.assertEqual(self.client.post(f"/api/review/tasks/{task_id}/comments", json={"comment": "补充核验交付验收条款"}).status_code, 200)
        self.assertEqual(self.client.post(f"/api/review/tasks/{task_id}/writeback").json()["status"], "success")
        report = self.client.get(f"/api/review/tasks/{task_id}/report/markdown")
        self.assertEqual(report.status_code, 200)
        self.assertIn("补充核验交付验收条款", report.text)
        self.assertEqual(self.client.get(f"/api/review/tasks/{task_id}/writeback").json()["status"], "success")
        pdf = self.client.get(f"/api/review/tasks/{task_id}/report/pdf")
        self.assertEqual(pdf.status_code, 200, pdf.text[:200])
        self.assertTrue(pdf.content.startswith(b"%PDF"))

    def test_demo_contains_expected_high_risks(self):
        result = self.client.post("/api/demo").json()
        self.assertEqual(result["task"]["status"], "completed")
        high_titles = {risk["title"] for risk in result["risks"] if risk["level"] == "HIGH"}
        self.assertIn("知识产权归属", high_titles)
        self.assertTrue(any(risk["clause_type"] == "payment" and risk["level"] == "HIGH" for risk in result["risks"]))
        self.assertTrue(any(risk["clause_type"] == "liability" and risk["level"] == "HIGH" for risk in result["risks"]))

    def test_workbench_assets_are_served(self):
        page = self.client.get("/")
        self.assertEqual(page.status_code, 200)
        self.assertIn("合同审查", page.text)
        for asset in ("app.js", "style.css"):
            response = self.client.get(f"/static/{asset}")
            self.assertEqual(response.status_code, 200)

    def test_normal_demo_and_role_permissions(self):
        normal = self.client.post("/api/demo?kind=normal").json()
        self.assertEqual(normal["risks"], [])
        business = self.client.post("/api/auth/login", json={"username": "business", "password": "Demo123!"}).json()
        headers = {"Authorization": f"Bearer {business['token']}"}
        self.assertEqual(self.client.get("/api/review/tasks", headers=headers).json(), [])
        self.assertEqual(self.client.get(f"/api/review/tasks/{normal['task']['id']}", headers=headers).status_code, 404)
        own = self.client.post("/api/demo?kind=normal", headers=headers).json()
        self.assertEqual(self.client.get("/api/review/tasks", headers=headers).json()[0]["name"], "软件采购合同正常示例")
        self.assertEqual(self.client.post(f"/api/review/tasks/{own['task']['id']}/writeback", headers=headers).status_code, 403)
        self.assertEqual(self.client.get("/api/review/tasks", headers={"Authorization": ""}).status_code, 401)

    def test_pdf_location_and_blocked_retry(self):
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont
        from reportlab.pdfgen import canvas

        pdfmetrics.registerFont(TTFont("ContractTestCJK", "C:/Windows/Fonts/msyh.ttc", subfontIndex=0))
        source = io.BytesIO()
        pdf = canvas.Canvas(source)
        pdf.setFont("ContractTestCJK", 11)
        pdf.drawString(72, 760, "采购合同第一页")
        pdf.showPage()
        pdf.setFont("ContractTestCJK", 11)
        pdf.drawString(72, 760, "项目开发成果及相关知识产权均归乙方所有。")
        pdf.save()
        uploaded = self.client.post("/api/contracts/upload", files={"file": ("risk.pdf", source.getvalue(), "application/pdf")}).json()
        task_id = uploaded["task"]["id"]
        self.uploads.append(storage.UPLOADS / f"{task_id}.pdf")
        ip_risk = next(risk for risk in uploaded["risks"] if risk["clause_type"] == "ip")
        self.assertEqual(ip_risk["page"], 2)

        admin = self.client.post("/api/auth/login", json={"username": "admin", "password": "Demo123!"}).json()
        headers = {"Authorization": f"Bearer {admin['token']}"}
        blocked = self.client.post("/api/contracts/upload", files={"file": ("scan.png", b"not-an-image", "image/png")}, headers=headers).json()
        self.uploads.append(storage.UPLOADS / f"{blocked['task']['id']}.png")
        self.assertEqual(blocked["task"]["status"], "blocked")
        self.assertTrue(blocked["task"]["blocked_reason"])
        retried = self.client.post(f"/api/review/tasks/{blocked['task']['id']}/retry", headers=headers).json()
        self.assertEqual(retried["task"]["status"], "blocked")
        self.assertTrue(retried["task"]["blocked_reason"])


if __name__ == "__main__":
    unittest.main()
